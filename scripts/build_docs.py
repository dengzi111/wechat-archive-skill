#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 messages.json + 原图 + 语音 生成 Word / HTML / CSV 成品。

用法:
    python build_docs.py --messages export/messages.json --out <输出目录> \
        [--images "<微信保存目录>"] [--voices export/voices] [--title "与 XX 的聊天记录"]

产出:
    <输出目录>/聊天记录.docx   图片以缩略图内嵌在对应消息位置（几百张也只有几 MB）
    <输出目录>/聊天记录.html   气泡排版 + 语音 <audio> + 图片点击放大（引用 图片/ 原图）
    <输出目录>/聊天记录.csv
    <输出目录>/图片/           原图副本（供 HTML 放大查看）
    <输出目录>/语音/           mp3 副本

依赖: pip install python-docx Pillow
"""
import argparse
import csv
import datetime
import html
import json
import os
import re
import shutil

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image

WEEK = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]


def build(args):
    msgs = json.load(open(args.messages, encoding="utf-8"))
    out = args.out
    imgdir = os.path.join(out, "图片")
    voicedir = os.path.join(out, "语音")
    thumbs = os.path.join(out, "_thumbs")
    for d in (out, imgdir, voicedir, thumbs):
        os.makedirs(d, exist_ok=True)

    # 文件名 -> 消息： 微信图片_<14位时间>_<local_id>_
    apat = re.compile(r"微信图片_(\d{14})_(\d+)_")
    vpat = re.compile(r"微信视频.*?_(\d+)\.")
    saved = {}
    if args.images and os.path.isdir(args.images):
        for f in os.listdir(args.images):
            m = apat.match(f)
            if m:
                saved[("3", m.group(1), int(m.group(2)))] = f
                continue
            if f.startswith("微信视频"):
                v = vpat.match(f)
                if v:
                    saved[("43", "", int(v.group(1)))] = f

    def stamp(t):
        return t[:4] + t[5:7] + t[8:10] + t[11:13] + t[14:16] + t[17:19]

    recs = []
    for m in msgs:
        t = m["type"]
        who = "我" if m.get("is_me") else "对方"
        st = stamp(m.get("time", ""))
        img = ""
        if t == 3:
            img = saved.get(("3", st, m["local_id"]), "")
        elif t == 43:
            img = saved.get(("43", "", m["local_id"]), "") or saved.get(("43", st, m["local_id"]), "")
        voice = m.get("voice", "")
        recs.append({"time": m.get("time", ""), "date": (m.get("time") or "")[:10], "who": who,
                     "type": t, "text": m.get("text", ""), "img": img, "voice": voice})

    for r in recs:
        if r["img"]:
            src = os.path.join(args.images, r["img"])
            if os.path.exists(src) and not os.path.exists(os.path.join(imgdir, r["img"])):
                shutil.copy2(src, os.path.join(imgdir, r["img"]))
    if args.voices and os.path.isdir(args.voices):
        for f in os.listdir(args.voices):
            d = os.path.join(voicedir, f)
            if not os.path.exists(d):
                shutil.copy2(os.path.join(args.voices, f), d)

    def thumb(name):
        dst = os.path.join(thumbs, os.path.splitext(name)[0] + ".jpg")
        if os.path.exists(dst):
            return dst
        try:
            im = Image.open(os.path.join(args.images, name)).convert("RGB")
            im.thumbnail((330, 330))
            im.save(dst, "JPEG", quality=68)
            return dst
        except Exception:
            return None

    title = args.title or "聊天记录"
    # ---- Word ----
    doc = Document()
    doc.styles["Normal"].font.name = "微软雅黑"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"消息总数：{len(recs)} 条")
    doc.add_paragraph(f"时间范围：{recs[0]['time']} ～ {recs[-1]['time']}" if recs else "")
    doc.add_paragraph(f"图片 {sum(1 for r in recs if r['img'])} 张已内嵌；语音 mp3 见「语音」文件夹；原图见「图片」文件夹。")
    doc.add_page_break()
    cur = None
    for r in recs:
        if r["date"] != cur:
            cur = r["date"]
            try:
                d = datetime.date.fromisoformat(cur)
                doc.add_heading(f"{d.year}年{d.month}月{d.day}日 {WEEK[d.weekday()]}", level=2)
            except Exception:
                doc.add_heading(cur, level=2)
        p = doc.add_paragraph()
        p.add_run(f"{r['time'][11:]}  {r['who']}：{r['text']}")
        if r["img"] and r["type"] == 3:
            t = thumb(r["img"])
            if t:
                try:
                    p.add_run().add_picture(t, width=Inches(2.1))
                except Exception:
                    pass
        if r["voice"]:
            run = p.add_run(f"  （音频：{r['voice']}）")
            run.italic = True
            run.font.size = Pt(9)
    docx_path = os.path.join(out, "聊天记录.docx")
    doc.save(docx_path)

    # ---- HTML ----
    hp = os.path.join(out, "聊天记录.html")
    with open(hp, "w", encoding="utf-8") as f:
        f.write("""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<title>""" + html.escape(title) + """</title><style>
body{background:#f2f2f7;font-family:"Microsoft YaHei",system-ui;margin:0;padding:16px 8px}
.wrap{max-width:980px;margin:0 auto}h1{font-size:20px;text-align:center}
.meta{text-align:center;color:#666;font-size:13px;margin-bottom:18px;line-height:1.7}
.day{text-align:center;margin:22px 0 12px;color:#888;font-size:13px}
.row{display:flex;margin:6px 0}.row.me{justify-content:flex-end}
.bubble{max-width:70%;padding:8px 12px;border-radius:12px;background:#fff;box-shadow:0 1px 2px rgba(0,0,0,.08);font-size:14px;line-height:1.5;white-space:pre-wrap;word-break:break-word}
.row.me .bubble{background:#95ec69}
.t{font-size:11px;color:#999;margin:0 6px;align-self:flex-end}
.audio{vertical-align:middle;height:32px}
img.ph{max-width:230px;border-radius:8px;display:block;margin-top:6px;cursor:zoom-in}
</style></head><body><div class="wrap">""")
        f.write(f"<h1>{html.escape(title)}</h1>")
        f.write(f"<div class='meta'>共 {len(recs)} 条 · {recs[0]['time']} ～ {recs[-1]['time']}</div>" if recs else "")
        cur = None
        for r in recs:
            if r["date"] != cur:
                cur = r["date"]
                f.write(f"<div class='day'>{cur}</div>")
            cls = "row me" if r["who"] == "我" else "row"
            f.write(f"<div class='{cls}'><div class='bubble'>{html.escape(r['text'])}")
            if r["img"] and r["type"] == 3:
                f.write(f"<img class='ph' loading='lazy' src='图片/{html.escape(r['img'])}' onclick=\"window.open(this.src)\">")
            if r["voice"]:
                f.write(f"<br><audio class='audio' controls preload='none' src='语音/{html.escape(r['voice'])}'></audio>")
            f.write(f"</div><div class='t'>{r['time'][11:16]}</div></div>")
        f.write("</div></body></html>")

    # ---- CSV ----
    with open(os.path.join(out, "聊天记录.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["时间", "发送人", "类型", "内容", "图片/视频", "语音"])
        for r in recs:
            w.writerow([r["time"], r["who"], r["type"], r["text"], r["img"], r["voice"]])

    print(f"完成: {out}\n  docx {os.path.getsize(docx_path)/1048576:.1f} MB, 图片副本 {len(os.listdir(imgdir))}, 语音副本 {len(os.listdir(voicedir))}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--messages", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--images", help="微信「另存为」保存目录")
    ap.add_argument("--voices", help="语音 mp3 目录")
    ap.add_argument("--title", default="聊天记录")
    build(ap.parse_args())


if __name__ == "__main__":
    main()
